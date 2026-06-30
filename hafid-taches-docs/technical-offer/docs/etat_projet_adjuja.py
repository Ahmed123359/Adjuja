# -- coding: utf-8 --
"""
Génération du document d'état du projet ADJUJA : description fonctionnelle et technique.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


FONT_NAME   = "Times New Roman"
COLOR_H1    = RGBColor(0x1F, 0x38, 0x64)
COLOR_H2    = RGBColor(0x2C, 0x3E, 0x50)
COLOR_BODY  = RGBColor(0x00, 0x00, 0x00)
COLOR_NOTE  = RGBColor(0x55, 0x55, 0x55)
SIZE_BODY   = Pt(11)
SIZE_H1     = Pt(13)
SIZE_H2     = Pt(12)


def _force_tnr(run, size: Pt) -> None:
    rpr = run._r.get_or_add_rPr()
    for existing in rpr.findall(qn("w:rFonts")):
        rpr.remove(existing)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT_NAME)
    rfonts.set(qn("w:hAnsi"), FONT_NAME)
    rfonts.set(qn("w:cs"),    FONT_NAME)
    for attr in (qn("w:asciiTheme"), qn("w:hAnsiTheme"), qn("w:cstheme"), qn("w:eastAsiaTheme")):
        if attr in rfonts.attrib:
            del rfonts.attrib[attr]
    rpr.insert(0, rfonts)
    run.font.size = size


def _set_doc_defaults(doc: Document) -> None:
    styles_el = doc.styles.element
    doc_defaults = styles_el.find(qn("w:docDefaults"))
    if doc_defaults is None:
        return
    rpr_default = doc_defaults.find(f".//{qn('w:rPrDefault')}/{qn('w:rPr')}")
    if rpr_default is None:
        return
    for existing in rpr_default.findall(qn("w:rFonts")):
        rpr_default.remove(existing)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT_NAME)
    rfonts.set(qn("w:hAnsi"), FONT_NAME)
    rfonts.set(qn("w:cs"),    FONT_NAME)
    rpr_default.insert(0, rfonts)


def para_format(para, space_after: int = 6, line_spacing: float = 1.15) -> None:
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = para.paragraph_format
    fmt.space_after  = Pt(space_after)
    fmt.line_spacing = Pt(11 * line_spacing)


def add_h1(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    para_format(p, space_after=8)
    p.paragraph_format.space_before = Pt(14)
    run = p.add_run(text)
    run.bold = True
    run.font.color.rgb = COLOR_H1
    _force_tnr(run, SIZE_H1)


def add_h2(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    para_format(p, space_after=6)
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run(text)
    run.bold = True
    run.font.color.rgb = COLOR_H2
    _force_tnr(run, SIZE_H2)


def add_body(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    para_format(p)
    run = p.add_run(text)
    run.bold = bold
    run.font.color.rgb = COLOR_BODY
    _force_tnr(run, SIZE_BODY)


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    para_format(p, space_after=4)
    run = p.add_run(text)
    run.font.color.rgb = COLOR_NOTE
    run.font.italic = True
    _force_tnr(run, Pt(10))


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    para_format(p, space_after=3)
    p.paragraph_format.left_indent = Pt(20)
    run = p.add_run(f"- {text}")
    run.font.color.rgb = COLOR_BODY
    _force_tnr(run, SIZE_BODY)


def add_code(doc: Document, lines: list) -> None:
    for line in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after  = Pt(0)
        p.paragraph_format.left_indent  = Pt(20)
        p.paragraph_format.right_indent = Pt(20)
        run = p.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        # Fond gris
        rpr = run._r.get_or_add_rPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"),   "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"),  "F2F2F2")
        rpr.append(shd)
        # Bordure gauche bleue marine sur le paragraphe
        ppr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"),   "single")
        left.set(qn("w:sz"),    "12")
        left.set(qn("w:space"), "4")
        left.set(qn("w:color"), "1F3864")
        pBdr.append(left)
        ppr.append(pBdr)
    doc.add_paragraph()


def add_table(doc: Document, headers: list, rows: list) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        cell = hdr_cells[i]
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = para.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _force_tnr(run, Pt(10))
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"),   "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"),  "1F3864")
        tc_pr.append(shd)
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        fill = "F2F2F2" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row_cells[c_idx]
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = para.add_run(str(val))
            run.font.color.rgb = COLOR_BODY
            _force_tnr(run, Pt(10))
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"),   "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"),  fill)
            tc_pr.append(shd)
    try:
        tbl_pr = table._tbl.get_or_add_tblPr()
        tbl_borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{side}")
            border.set(qn("w:val"),   "single")
            border.set(qn("w:sz"),    "4")
            border.set(qn("w:space"), "0")
            border.set(qn("w:color"), "BFBFBF")
            tbl_borders.append(border)
        tbl_pr.append(tbl_borders)
    except Exception:
        pass
    doc.add_paragraph()


def build() -> Document:
    doc = Document()
    _set_doc_defaults(doc)

    section = doc.sections[0]
    section.page_width    = Inches(8.27)
    section.page_height   = Inches(11.69)
    section.left_margin   = Inches(1.18)
    section.right_margin  = Inches(1.18)
    section.top_margin    = Inches(1.18)
    section.bottom_margin = Inches(1.0)

    # ── Page de titre ─────────────────────────────────────────────────────────
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after  = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("ADJUJA")
    r.bold = True
    r.font.color.rgb = COLOR_H1
    _force_tnr(r, Pt(20))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Etat du projet : description fonctionnelle et technique")
    r.bold = True
    r.font.color.rgb = COLOR_H2
    _force_tnr(r, Pt(14))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(60)
    r = p.add_run("Plateforme SaaS de génération automatisée de réponses aux appels d'offres publics")
    r.font.color.rgb = COLOR_NOTE
    _force_tnr(r, Pt(12))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Mai 2026")
    r.font.color.rgb = COLOR_NOTE
    _force_tnr(r, SIZE_BODY)

    doc.add_page_break()

    # ── Table des matières ────────────────────────────────────────────────────
    add_h1(doc, "Table des matières")
    toc = [
        ("PARTIE I.", "Description fonctionnelle"),
        ("1.", "Vue d'ensemble de la plateforme"),
        ("1.1.", "Positionnement et cible"),
        ("1.2.", "Architecture de l'interface"),
        ("2.", "Les pages de l'application"),
        ("2.1.", "Page d'accueil publique"),
        ("2.2.", "Authentification"),
        ("2.3.", "Tableau de bord"),
        ("2.4.", "Appels d'offres"),
        ("2.5.", "Outils"),
        ("2.6.", "Paramètres"),
        ("2.7.", "Chatbot intégré"),
        ("3.", "Parcours de génération d'un dossier AO"),
        ("3.1.", "Préparation du profil entreprise"),
        ("3.2.", "Création d'un appel d'offres"),
        ("3.3.", "Lancement du pipeline"),
        ("3.4.", "Suivi et téléchargement"),
        ("PARTIE II.", "Description technique"),
        ("4.", "Stack technologique"),
        ("4.1.", "Backend"),
        ("4.2.", "Frontend"),
        ("4.3.", "Infrastructure"),
        ("5.", "Architecture du pipeline de génération"),
        ("5.1.", "Vue d'ensemble des tâches Celery"),
        ("5.2.", "Classification et analyse du dossier AO"),
        ("5.3.", "Génération de l'offre technique"),
        ("5.4.", "Remplissage des documents administratifs"),
        ("5.5.", "Matching de l'équipe"),
        ("5.6.", "Signature et compilation"),
        ("6.", "Base de connaissances RAG"),
        ("7.", "Système de signature des documents"),
        ("8.", "Modèles de données"),
    ]
    for num, title in toc:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{num}   {title}")
        r.font.color.rgb = COLOR_BODY
        r.bold = ("." not in num[1:] or num.endswith(".") and len(num) <= 3)
        _force_tnr(r, SIZE_BODY)

    doc.add_page_break()

    # =========================================================================
    # PARTIE I  FONCTIONNELLE
    # =========================================================================
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("PARTIE I  :  DESCRIPTION FONCTIONNELLE")
    r.bold = True
    r.font.color.rgb = COLOR_H1
    _force_tnr(r, Pt(13))

    # 1. Vue d'ensemble
    add_h1(doc, "1. Vue d'ensemble de la plateforme")

    add_h2(doc, "1.1. Positionnement et cible")
    add_body(doc,
        "ADJUJA est une plateforme SaaS destinée aux cabinets de conseil et aux PME marocaines "
        "qui répondent régulièrement aux appels d'offres publics. L'objectif est de réduire le "
        "temps de préparation d'un dossier de soumission de plusieurs jours à moins d'une heure, "
        "en automatisant la génération des documents obligatoires : offre technique, acte "
        "d'engagement, déclaration sur l'honneur et bordereau des prix."
    )
    add_body(doc,
        "La plateforme fonctionne selon un principe de configuration unique : l'utilisateur "
        "renseigne son profil entreprise, upload ses CVs et ses documents permanents une seule "
        "fois. Ensuite, pour chaque nouvel appel d'offres, deux fichiers suffisent (CPS et RC) "
        "pour déclencher la génération complète du dossier."
    )

    add_h2(doc, "1.2. Architecture de l'interface")
    add_body(doc,
        "L'interface est organisée autour d'une barre latérale fixe de 208 pixels qui donne "
        "accès aux trois espaces principaux de l'application : les appels d'offres, les outils "
        "et les paramètres. Un bouton de bascule thème clair/sombre et un sélecteur de langue "
        "(français/anglais) sont présents en permanence dans la barre supérieure. L'ensemble "
        "de l'application respecte une largeur maximale de contenu de 896 pixels pour les pages "
        "de liste et 640 pixels pour les formulaires."
    )

    # 2. Pages
    add_h1(doc, "2. Les pages de l'application")

    add_h2(doc, "2.1. Page d'accueil publique")
    add_body(doc,
        "La page d'accueil présente la plateforme aux visiteurs non connectés. Elle contient "
        "une section héro avec un titre accrocheur, un mockup animé de l'interface et un bouton "
        "d'appel à l'action vers l'inscription. Deux sections suivent : une présentation des "
        "fonctionnalités principales et une description du processus en étapes. La navigation "
        "supérieure propose des liens vers ces sections et vers les pages de connexion et "
        "d'inscription."
    )

    add_h2(doc, "2.2. Authentification")
    add_body(doc,
        "La page d'inscription collecte le nom, le prénom, l'adresse électronique et le mot de "
        "passe. La connexion utilise l'adresse électronique et le mot de passe. L'authentification "
        "repose sur des tokens JWT. Toute page de l'application est protégée : un utilisateur non "
        "connecté est redirigé vers la page de connexion."
    )

    add_h2(doc, "2.3. Tableau de bord")
    add_body(doc,
        "Le tableau de bord est organisé en six onglets accessibles via une barre de navigation "
        "horizontale."
    )
    add_table(doc,
        headers=["Onglet", "Contenu"],
        rows=[
            ["Vue d'ensemble",      "Statistiques globales : nombre d'AOs, documents générés, statuts. Accès rapide aux derniers appels d'offres."],
            ["Profil entreprise",   "Formulaire complet des données de l'entreprise : raison sociale, ICE, RC, IF, CNSS, capital social, RIB, forme juridique, adresse, ville, téléphone, email, gérant et CIN."],
            ["Paraphe & Cachet",    "Upload de l'image du paraphe du gérant (utilisé sur CPS et RC). Upload du cachet de l'entreprise (apposé sur les documents administratifs). Upload optionnel d'une image 'Lu et accepté'. Upload optionnel d'un template DOCX pour la note méthodologique."],
            ["Documents",           "Gestion des documents permanents de l'entreprise : PV de gérance, attestations fiscales et CNSS, note sur les moyens. Ces documents sont inclus automatiquement dans chaque dossier généré."],
            ["Equipe / CVs",        "Gestion du pool de formateurs et experts. Upload d'un PDF de CV : le système extrait automatiquement les informations (nom, prénom, poste, spécialité, diplôme, expérience) via Mistral. Le CV est ensuite indexé dans la base RAG."],
            ["Génération",          "Interface de démarrage rapide pour créer un AO directement depuis le tableau de bord."],
        ]
    )

    add_h2(doc, "2.4. Appels d'offres")
    add_body(doc,
        "La page des appels d'offres liste tous les dossiers créés pour l'organisation. Chaque "
        "entrée affiche le nom de l'AO, son statut (en attente, en traitement, terminé, erreur) "
        "et un bouton de suppression. Un bouton 'Nouvel AO' permet de créer une nouvelle entrée."
    )
    add_body(doc,
        "Au clic sur un AO, le détail s'ouvre et présente : la barre de progression du pipeline "
        "(de 0 à 100 %), la liste des documents sources uploadés (CPS, RC) avec boutons de "
        "téléchargement, les documents générés organisés par dossier (financier, administratif, "
        "technique) avec boutons PDF et DOCX selon le format disponible, et le dossier complet "
        "en ZIP téléchargeable en un clic."
    )

    add_h2(doc, "2.5. Outils")
    add_body(doc,
        "La page Outils donne accès aux documents permanents de l'entreprise : consultation, "
        "upload et suppression. C'est une vue simplifiée du même contenu que l'onglet Documents "
        "du tableau de bord, accessible directement depuis la navigation principale."
    )

    add_h2(doc, "2.6. Paramètres")
    add_body(doc,
        "Les paramètres proposent deux onglets. Le premier, 'Général', permet de choisir le "
        "fournisseur LLM par défaut pour le chatbot parmi Mistral, OpenAI et Anthropic. Le "
        "second, 'Utilisation', affiche des indicateurs de consommation de l'API."
    )

    add_h2(doc, "2.7. Chatbot intégré")
    add_body(doc,
        "Un chatbot flottant est accessible depuis toutes les pages de l'application via une "
        "bulle en bas à droite de l'écran. Il répond aux questions de l'utilisateur en s'appuyant "
        "sur la base de connaissances interne de l'organisation (CVs, documents entreprise, "
        "analyses d'AOs passés). L'utilisateur peut choisir le modèle LLM utilisé pour la "
        "réponse. L'historique de la conversation est conservé pour la session en cours."
    )

    doc.add_page_break()

    # 3. Parcours
    add_h1(doc, "3. Parcours de génération d'un dossier AO")

    add_h2(doc, "3.1. Préparation du profil entreprise")
    add_body(doc,
        "Avant de traiter un premier AO, l'utilisateur complète le profil de son entreprise "
        "depuis le tableau de bord. Cette étape est réalisée une seule fois. Elle couvre "
        "l'identité juridique, les coordonnées, les données financières et les représentants "
        "légaux. L'utilisateur uploade ensuite le paraphe du gérant, le cachet de l'entreprise "
        "et les documents permanents (attestations, PV de gérance)."
    )
    add_body(doc,
        "Le pool de CVs est alimenté de la même façon : pour chaque expert ou formateur, "
        "l'utilisateur dépose le PDF du CV. Le système extrait les informations en quelques "
        "secondes et présente un formulaire pré-rempli à valider. Une fois sauvegardé, le CV "
        "est indexé et disponible pour le matching automatique."
    )

    add_h2(doc, "3.2. Création d'un appel d'offres")
    add_body(doc,
        "L'utilisateur clique sur 'Nouvel AO' et saisit un nom identifiant le dossier. "
        "L'interface lui présente ensuite une zone d'upload pour le CPS (cahier des prescriptions "
        "spéciales) et le RC (règlement de consultation). Ces deux fichiers PDF sont les seuls "
        "documents source nécessaires au pipeline."
    )

    add_h2(doc, "3.3. Lancement du pipeline")
    add_body(doc,
        "Une fois les fichiers uploadés, l'utilisateur clique sur 'Démarrer'. Le pipeline se "
        "lance en tâche de fond. La barre de progression se met à jour en temps réel, par "
        "interrogation périodique du serveur. L'utilisateur n'a rien d'autre à faire."
    )
    add_body(doc,
        "Le pipeline produit automatiquement : une note méthodologique structurée en dix "
        "sections adaptées au contenu du CPS, un acte d'engagement pré-rempli avec les données "
        "de l'entreprise, une déclaration sur l'honneur pré-remplie, un bordereau des prix "
        "formaté, et l'ensemble des documents permanents. Chaque document PDF est signé avec "
        "le cachet ou le paraphe selon sa nature."
    )

    add_h2(doc, "3.4. Suivi et téléchargement")
    add_body(doc,
        "Dès que le pipeline atteint 100 %, le statut passe à 'Terminé' et la liste des "
        "documents apparaît. L'utilisateur peut télécharger chaque document individuellement "
        "au format PDF ou DOCX, ou récupérer le dossier complet en un fichier ZIP. Les "
        "documents administratifs et technique sont rangés dans des sous-dossiers distincts "
        "à l'intérieur du ZIP."
    )

    doc.add_page_break()

    # =========================================================================
    # PARTIE II  TECHNIQUE
    # =========================================================================
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("PARTIE II  :  DESCRIPTION TECHNIQUE")
    r.bold = True
    r.font.color.rgb = COLOR_H1
    _force_tnr(r, Pt(13))

    # 4. Stack
    add_h1(doc, "4. Stack technologique")

    add_h2(doc, "4.1. Backend")
    add_body(doc,
        "Le backend est une API REST construite avec FastAPI et Python. La base de données "
        "est PostgreSQL, accédée de façon asynchrone via SQLAlchemy et asyncpg. Les migrations "
        "de schéma sont gérées par Alembic. Les fichiers binaires (PDFs, images, documents) "
        "sont stockés dans MinIO, un serveur objet compatible S3. Redis sert à la fois de "
        "broker de messages et de backend de résultats pour Celery. Qdrant est la base "
        "vectorielle utilisée pour la recherche RAG."
    )
    add_table(doc,
        headers=["Composant", "Rôle"],
        rows=[
            ["FastAPI",      "API REST asynchrone, gestion des routes, validation des données avec Pydantic"],
            ["PostgreSQL",   "Base de données relationnelle : utilisateurs, organisations, AOs, documents, profils"],
            ["MinIO",        "Stockage objet des fichiers binaires : PDFs, images, DOCX générés"],
            ["Redis",        "Broker Celery (base 1), backend de résultats (base 2), cache général (base 0)"],
            ["Celery",       "Orchestrateur de tâches asynchrones, deux workers : celery_io et celery_cpu"],
            ["Qdrant",       "Base vectorielle pour la recherche sémantique (RAG) sur les documents de l'organisation"],
            ["Mistral AI",   "Fournisseur LLM principal : génération de texte, extraction JSON, embeddings"],
        ]
    )

    add_h2(doc, "4.2. Frontend")
    add_body(doc,
        "Le frontend est une application React avec TypeScript, compilée par Vite. Le style "
        "est géré par une combinaison de CSS custom properties (variables --l-*) et de Tailwind "
        "pour les utilitaires de mise en page. L'internationalisation est assurée par "
        "react-i18next avec deux fichiers de traduction (français et anglais). La police "
        "principale est DM Sans. Tous les appels à l'API passent par un module centralisé "
        "api.ts qui injecte automatiquement le token d'authentification."
    )

    add_h2(doc, "4.3. Infrastructure")
    add_body(doc,
        "En développement, tous les services sont orchestrés par Docker Compose. L'API et "
        "le frontend tournent localement avec hot-reload. En production, le même fichier "
        "Docker Compose sans suffixe orchestre les conteneurs, et Nginx sert le frontend "
        "compilé ainsi que le proxy vers l'API."
    )
    add_body(doc,
        "Deux workers Celery distincts traitent les tâches selon leur nature. Le worker "
        "celery_io traite les tâches d'entrée/sortie légères : classification, analyse, "
        "matching d'équipe, indexation. Le worker celery_cpu traite les tâches intensives : "
        "génération de la note méthodologique, remplissage de documents, signature et "
        "compilation du ZIP."
    )

    doc.add_page_break()

    # 5. Pipeline
    add_h1(doc, "5. Architecture du pipeline de génération")

    add_h2(doc, "5.1. Vue d'ensemble des tâches Celery")
    add_body(doc,
        "Le pipeline suit un graphe de tâches précis. Une chaîne séquentielle de trois "
        "tâches initialise l'analyse. A l'issue de cette chaîne, un accord (chord) Celery "
        "lance en parallèle la génération de l'offre technique et le remplissage des "
        "documents administratifs. La signature et la compilation du ZIP n'interviennent "
        "que lorsque les deux branches parallèles sont terminées. Le matching de l'équipe "
        "s'exécute en parallèle de l'accord, sans bloquer le pipeline."
    )
    add_code(doc, [
        "chain(",
        "    task_classify_uploads(ao_id),",
        "    task_analyze_ao_context(ao_id),",
        "    task_build_pipeline(ao_id),",
        ")",
        "chord(",
        "    group(",
        "        task_generate_note_metho(ao_id),",
        "        task_fill_documents(ao_id),",
        "    ),",
        "    task_sign_and_compile(ao_id),",
        ")",
        "task_match_team(ao_id)  # parallele, queue celery_io",
        "task_index_results(ao_id)",
    ])

    add_h2(doc, "5.2. Classification et analyse du dossier AO")
    add_body(doc,
        "La première tâche, task_classify_uploads, détermine le type de chaque document "
        "uploadé (CPS ou RC) par analyse du nom de fichier et du contenu textuel extrait "
        "par PyMuPDF. La tâche suivante, task_analyze_ao_context, envoie le texte complet "
        "du CPS et du RC à Mistral Large. Le modèle retourne un objet JSON structuré "
        "contenant : le contexte de l'appel d'offres (acheteur, objet, lots, délai), la "
        "liste des profils requis avec diplôme et expérience minimum, la liste des documents "
        "à produire avec leur source (générer, remplir ou signer), et les critères de "
        "pondération de l'offre technique."
    )
    add_body(doc,
        "La tâche task_build_pipeline lit cet objet JSON pour décider si la génération "
        "de l'offre technique est nécessaire et quels templates de documents doivent être "
        "remplis. Elle configure et lance l'accord Celery."
    )

    add_h2(doc, "5.3. Génération de l'offre technique")
    add_body(doc,
        "La tâche task_generate_note_metho est la plus longue du pipeline, avec une durée "
        "typique de deux à trois minutes. Elle délègue à un service dédié qui exécute "
        "plusieurs sous-étapes."
    )
    add_body(doc,
        "Le service commence par analyser le CPS avec un extracteur spécialisé qui identifie "
        "le périmètre, les livrables, le vocabulaire propre au maître d'ouvrage et les "
        "contraintes de délai. Le RC est analysé séparément pour extraire les critères "
        "d'évaluation et les exigences de format. Un moteur de stratégie choisit ensuite "
        "l'angle de l'offre selon le type de mission."
    )
    add_body(doc,
        "La génération des sections utilise six appels Mistral Large en parallèle, limités "
        "à deux appels simultanés pour respecter les limites de débit de l'API. Chaque "
        "appel reçoit un prompt système spécifique à sa section, le contexte extrait du "
        "CPS, le contexte RAG récupéré depuis Qdrant, et les informations sur l'équipe "
        "matchée. Un contrôle qualité évalue le résultat et peut déclencher une régénération "
        "de la section la plus faible, jusqu'à deux fois."
    )
    add_body(doc,
        "Une fois les sections générées, l'assembleur DOCX construit le document final en "
        "dix sections : page de garde, sommaire, présentation du cabinet, compréhension du "
        "contexte, compréhension de la mission, approche méthodologique, équipe proposée, "
        "planning d'exécution avec diagramme Gantt, chronogramme d'affectation du personnel "
        "et annexes. Le document est ensuite converti en PDF par LibreOffice."
    )

    add_h2(doc, "5.4. Remplissage des documents administratifs")
    add_body(doc,
        "La tâche task_fill_documents parcourt la liste des documents à remplir identifiés "
        "lors de l'analyse. Pour chaque document (acte d'engagement, déclaration sur "
        "l'honneur, bordereau des prix), elle appelle le service filler."
    )
    add_body(doc,
        "Le filler extrait le texte du PDF page par page, identifie les zones blanches et "
        "les champs à compléter, puis envoie ces éléments à Mistral Large avec les données "
        "de l'entreprise (company_info). Le modèle retourne un dictionnaire associant chaque "
        "identifiant de champ à sa valeur. Le filler reconstruit ensuite le PDF avec les "
        "valeurs injectées, en préservant la mise en page originale. Si le document est "
        "scanné (texte non extractible), Pixtral analyse les images page par page."
    )

    add_h2(doc, "5.5. Matching de l'équipe")
    add_body(doc,
        "La tâche task_match_team s'exécute en parallèle du reste du pipeline. Elle lit "
        "la liste des profils requis depuis l'objet analyse_json stocké en base et le pool "
        "de CVs actifs de l'organisation. Elle envoie les deux listes à Mistral Large qui "
        "retourne un tableau d'associations : chaque profil requis est associé au CV le plus "
        "adapté, avec un indicateur d'alerte si aucun CV ne correspond aux critères minimum. "
        "Les associations sont stockées dans la table ao_team_members."
    )
    add_body(doc,
        "La génération de la note méthodologique attend jusqu'à 90 secondes la complétion "
        "du matching avant de construire la section 'Equipe proposée'. En pratique, le "
        "matching se termine en moins de 15 secondes, avant que les premières sections "
        "LLM de l'offre technique ne soient produites."
    )

    add_h2(doc, "5.6. Signature et compilation")
    add_body(doc,
        "La tâche task_sign_and_compile est déclenchée automatiquement par Celery lorsque "
        "les deux branches parallèles (offre technique et filler) sont terminées. Elle "
        "charge les images du paraphe, du cachet et du tampon 'Lu et accepté' depuis MinIO."
    )
    add_body(doc,
        "Pour chaque document PDF, le service de signature applique un traitement différencié "
        "selon la nature du document. Les documents source CPS et RC reçoivent le paraphe "
        "du gérant sur chaque page et le tampon 'Lu et accepté' sur la dernière page. Les "
        "documents administratifs (acte d'engagement, déclaration sur l'honneur) et l'offre "
        "technique reçoivent uniquement le cachet de l'entreprise dans la zone de signature "
        "détectée automatiquement (zone 'Fait à' ou zone de mots-clés sur la dernière page). "
        "Le texte 'Fait à [ville], le [date]' est injecté dans les documents qui le prévoient."
    )
    add_body(doc,
        "Chaque PDF signé est uploadé en MinIO sous un chemin dédié et le minio_key du "
        "document en base de données est mis à jour pour pointer vers la version signée. "
        "Les boutons de téléchargement de l'interface servent donc directement la version "
        "signée. Un fichier ZIP est créé avec tous les documents organisés en dossiers "
        "(source, financier, administratif, technique) et uploadé en MinIO. Le statut du "
        "pipeline passe à 100 % et l'AO passe en statut 'terminé'."
    )

    doc.add_page_break()

    # 6. RAG
    add_h1(doc, "6. Base de connaissances RAG")
    add_body(doc,
        "ADJUJA maintient une base vectorielle par organisation dans Qdrant. La collection "
        "est nommée selon le schéma offria_kb_{org_id}. Chaque document indexé est découpé "
        "en chunks de 500 tokens avec un chevauchement de 50 tokens. Les embeddings sont "
        "calculés par le modèle Mistral Embed."
    )
    add_body(doc,
        "Trois types de documents alimentent la base. Les CVs des experts sont indexés "
        "automatiquement lors de leur upload depuis le tableau de bord. Les documents "
        "permanents de l'entreprise (attestations de référence, note sur les moyens) sont "
        "indexés à leur upload. Les analyses d'AOs terminés (objet, contexte, approche) "
        "sont indexées par la tâche task_index_results à la fin de chaque pipeline réussi."
    )
    add_body(doc,
        "Lors de la génération de l'offre technique, le service RAG est interrogé pour "
        "chacune des six sections avec une requête construite à partir du titre de la section "
        "et du périmètre de l'AO. Le contexte retourné est injecté dans le prompt de "
        "génération. Le chatbot utilise le même service RAG pour contextualiser ses réponses "
        "avec les documents de l'organisation."
    )

    # 7. Signature
    add_h1(doc, "7. Système de signature des documents")
    add_body(doc,
        "Le service de signature utilise PyMuPDF pour modifier les PDFs en mémoire sans "
        "passer par un outil externe. Pour chaque page, une série de détections est "
        "exécutée. La détection de la zone 'Fait à' cherche le patron textuel dans les "
        "spans de la page. La détection de zone de signature cherche des mots-clés comme "
        "'signature du soumissionnaire' ou 'cachet du concurrent' sur la dernière page "
        "uniquement."
    )

    add_table(doc,
        headers=["Document", "Paraphe (chaque page)", "Zone de signature (dernière page)", "Lu et accepté"],
        rows=[
            ["CPS",                   "Oui (image paraphe du gérant)", "Lu et accepté",                            "Oui (dernière page)"],
            ["RC",                    "Oui (image paraphe du gérant)", "Lu et accepté",                            "Oui (dernière page)"],
            ["Acte d'engagement",     "Non",                           "Cachet de l'entreprise uniquement",        "Non"],
            ["Déclaration honneur",   "Non",                           "Cachet de l'entreprise uniquement",        "Non"],
            ["Note méthodologique",   "Non",                           "Cachet de l'entreprise uniquement",        "Non"],
        ]
    )
    add_note(doc,
        "Le cachet de l'entreprise contient visuellement la signature du gérant. "
        "Il n'est donc pas nécessaire d'apposer l'image paraphe séparément sur les "
        "documents administratifs."
    )

    # 8. Modèles
    add_h1(doc, "8. Modèles de données")
    add_body(doc,
        "Les principales tables de la base de données sont les suivantes."
    )
    add_table(doc,
        headers=["Table", "Description"],
        rows=[
            ["organisations",       "Entités clientes, une par société"],
            ["users",               "Utilisateurs rattachés à une organisation"],
            ["company_profiles",    "Profil complet de l'entreprise : données légales, financières, clés MinIO des images et templates"],
            ["company_documents",   "Documents permanents (attestations, PV) avec clé MinIO et type"],
            ["staff_cvs",           "CVs des experts : données extraites, spécialité, statut actif"],
            ["appels_offres",       "Entrée par AO : statut, pourcentage pipeline, analyse JSON, instructions personnalisées"],
            ["ao_documents",        "Fichiers associés à un AO : type, dossier, origine, clé MinIO, statut"],
            ["ao_team_members",     "Associations AO vers staff_cv issues du matching, avec rôle et indicateur d'alerte"],
        ]
    )
    add_body(doc,
        "Le champ org_id est présent sur toutes les tables métier et vaut, en développement "
        "mono-tenant, l'identifiant de l'utilisateur lui-même. En configuration multi-tenant, "
        "il correspond à l'identifiant de l'organisation."
    )

    return doc


if __name__ == "__main__":
    output = "/home/almukhadram/Desktop/offria/reponse-ao-generation/hafid-taches-docs/technical-offer/docs/etat_projet_adjuja.docx"
    doc = build()
    doc.save(output)
    print(f"Document généré : {output}")
