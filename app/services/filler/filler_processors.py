#!/usr/bin/env python3
"""
filler_processors.py
--------------------
Pipelines de traitement des documents et fonctions de détection.

Pipelines exposés :
    - process_text_pdf()    : PDF contenant du texte extractible
    - process_docx()        : fichiers Word (.docx)
    - process_scanned_pdf() : PDF scanné (images uniquement)

Fonctions de détection :
    - is_scanned_pdf()       : distingue PDF texte et PDF scanné
    - detect_document_type() : identifie la catégorie du document
"""

import re
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

try:
    import fitz
except ImportError:
    sys.exit("Installer PyMuPDF : pip install PyMuPDF")

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

from app.services.filler.filler_llm import call_mistral, call_pixtral_vision
from app.services.filler.filler_placeholders import PH_RE, build_fill_operations, extract_replacement_pairs, has_placeholder
from app.services.filler.filler_settings import (
    DOCUMENT_TYPE_KEYWORDS,
    PAGE_HEIGHT_PT,
    PAGE_MARGIN_BOT_PT,
    PAGE_MARGIN_TOP_PT,
    PAGE_MARGIN_X_PT,
    PAGE_WIDTH_PT,
    RECONSTRUCTION_FONT,
    SCAN_DPI,
    SCANNED_IMAGE_AREA_RATIO,
    SCANNED_MAX_SUSPICIOUS_RATIO,
    SCANNED_MIN_READABLE_WORD_RATIO,
    SCANNED_MIN_TEXT_CHARS,
    SCANNED_MIN_WHITESPACE_RATIO,
)

SAFE_FONTS = {"helv", "Helvetica", "Times-Roman", "Courier", "Times", "Symbol", "ZapfDingbats"}


def _pdf_safe_text(text: str) -> str:
    replacements = {
        "’": "'", "‘": "'", "“": '"', "”": '"',
        "–": "-", "—": "-", "•": "-", " ": " ",
        "…": "...", "°": "°",
    }
    for src, dst in replacements.items():
        text = text.replace(src, dst)
    return text


# ===========================================================================
# DÉTECTION
# ===========================================================================

def is_scanned_pdf(pdf_path: Path) -> bool:
    doc       = fitz.open(str(pdf_path))
    full_text = "\n".join(page.get_text() for page in doc)

    for page in doc:
        page_area = page.rect.width * page.rect.height
        for image in page.get_images(full=True):
            for image_rect in page.get_image_rects(image[0]):
                image_area = image_rect.width * image_rect.height
                if page_area and image_area / page_area > SCANNED_IMAGE_AREA_RATIO:
                    doc.close()
                    return True

    doc.close()

    if len(full_text.strip()) < SCANNED_MIN_TEXT_CHARS:
        return True

    text_length    = max(len(full_text), 1)
    whitespace_chars = sum(c.isspace() for c in full_text)
    suspicious_chars = sum(
        not (c.isalnum() or c.isspace() or c in ".,;:!?()[]{}'\"/-_")
        for c in full_text
    )
    words          = re.findall(r"[A-Za-zÀ-ÿ]{3,}", full_text)
    readable_words = [w for w in words if re.search(r"[aeiouyAEIOUYàâäéèêëîïôöùûü]", w)]

    suspicious_ratio = suspicious_chars / text_length
    whitespace_ratio  = whitespace_chars  / text_length
    readable_ratio    = len(readable_words) / max(len(words), 1)

    return (
        suspicious_ratio > SCANNED_MAX_SUSPICIOUS_RATIO
        or whitespace_ratio  < SCANNED_MIN_WHITESPACE_RATIO
        or readable_ratio    < SCANNED_MIN_READABLE_WORD_RATIO
    )


def detect_document_type(text: str) -> str:
    normalized = text.lower()
    scores: dict[str, int] = {
        doc_type: sum(1 for kw in keywords if kw.lower() in normalized)
        for doc_type, keywords in DOCUMENT_TYPE_KEYWORDS.items()
    }
    best_type = max(scores, key=scores.get)
    return best_type if scores[best_type] > 0 else "unknown"


# ===========================================================================
# PIPELINE 1 : PDF TEXTE
# ===========================================================================

def _get_font_at(page: Any, rect: Any) -> tuple[str, float]:
    clip   = fitz.Rect(rect.x0 - 20, rect.y0 - 20, rect.x1 + 20, rect.y1 + 20)
    blocks = page.get_text("dict", clip=clip)["blocks"]
    for block in blocks:
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                if span.get("size"):
                    return span.get("font", "helv"), span["size"]
    return "helv", 10.0


def fill_text_pdf(src: Path, dst: Path, fill_ops: list[dict[str, Any]]) -> None:
    doc = fitz.open(str(src))
    operations_by_page: dict[int, list[dict[str, Any]]] = {}
    for op in fill_ops:
        operations_by_page.setdefault(op["page_no"], []).append(op)

    for page_no, operations in operations_by_page.items():
        page    = doc[page_no]
        pending: list[tuple[Any, str, str, float]] = []

        for op in operations:
            clip        = fitz.Rect(op["line_bbox"])
            search_clip = fitz.Rect(clip.x0, clip.y0 - 3, clip.x1 + 300, clip.y1 + 3)
            old_text    = op["old_text"]
            new_text    = op["new_text"]

            hits = page.search_for(old_text, clip=search_clip, quads=False)
            if not hits:
                all_hits = page.search_for(old_text, quads=False)
                if all_hits:
                    clip_cy = (clip.y0 + clip.y1) / 2
                    hits = [min(all_hits, key=lambda r: abs((r.y0 + r.y1) / 2 - clip_cy))]

            if not hits:
                print(f"  [WARN] Placeholder introuvable : {old_text[:30]!r} (page {page_no})")
                continue

            rect       = hits[0]
            font_name, font_size = _get_font_at(page, clip)
            safe_font  = font_name if font_name in SAFE_FONTS else "helv"
            pending.append((rect, new_text, safe_font, font_size))

        if not pending:
            continue

        for rect, _, _, _ in pending:
            page.add_redact_annot(rect, fill=(1, 1, 1))
        page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE)

        for rect, new_text, safe_font, font_size in pending:
            point = fitz.Point(rect.x0, rect.y1 - font_size * 0.15)
            page.insert_text(
                point, new_text,
                fontname=safe_font, fontsize=font_size, color=(0, 0, 0),
            )

    doc.save(str(dst), garbage=4, deflate=True)
    doc.close()


def process_text_pdf(
    src: Path,
    dst: Path,
    api_key: str,
    pages: list[int] | None = None,
    company_info: dict[str, str] | None = None,
) -> None:
    doc          = fitz.open(str(src))
    page_indices = pages if pages is not None else list(range(len(doc)))
    line_items:  list[dict[str, str]]        = []
    line_meta:   dict[str, dict[str, Any]]   = {}

    for page_no in page_indices:
        page   = doc[page_no]
        blocks = page.get_text("dict")["blocks"]
        for block in blocks:
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                spans = line.get("spans", [])
                text  = "".join(span["text"] for span in spans)
                if not has_placeholder(text):
                    continue
                line_id = f"p{page_no}_b{block['number']}_l{len(line_meta)}"
                line_meta[line_id] = {
                    "page_no":   page_no,
                    "bbox":      line["bbox"],
                    "font_name": spans[0].get("font", "helv") if spans else "helv",
                    "font_size": spans[0].get("size", 10.0)   if spans else 10.0,
                    "original":  text,
                }
                line_items.append({"id": line_id, "text": text})

    doc.close()

    if not line_items:
        print("  Aucun placeholder trouvé, copie sans modification.")
        shutil.copy(src, dst)
        return

    print(f"  {len(line_items)} ligne(s) avec placeholders trouvée(s).")
    print("  Envoi à Mistral pour remplissage...")
    filled_lines = call_mistral(line_items, api_key, company_info=company_info)
    print(f"  LLM a rempli {len(filled_lines)} ligne(s).")

    fill_ops = build_fill_operations(filled_lines, line_meta)
    if not fill_ops:
        print("  Aucun remplacement extrait — document inchangé.")
        shutil.copy(src, dst)
        return

    fill_text_pdf(src, dst, fill_ops)


# ===========================================================================
# PIPELINE 2 : DOCX
# ===========================================================================

def _replace_in_runs(runs: list[Any], old: str, new: str) -> bool:
    for run in runs:
        if old in run.text:
            run.text = run.text.replace(old, new, 1)
            return True
    full_text = "".join(run.text for run in runs)
    if old not in full_text:
        return False
    runs[0].text = full_text.replace(old, new, 1)
    for run in runs[1:]:
        run.text = ""
    return True


def process_docx(
    src: Path,
    dst: Path,
    api_key: str,
    company_info: dict[str, str] | None = None,
) -> None:
    if not DOCX_AVAILABLE:
        sys.exit("Installer python-docx : pip install python-docx")

    doc        = DocxDocument(str(src))
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)

    line_items = [
        {"id": str(i), "text": para.text}
        for i, para in enumerate(paragraphs)
        if has_placeholder(para.text)
    ]

    if not line_items:
        print("  Aucun placeholder trouvé.")
        shutil.copy(src, dst)
        return

    print(f"  {len(line_items)} paragraphe(s) avec placeholders trouvé(s).")
    filled = call_mistral(line_items, api_key, company_info=company_info)
    print(f"  LLM a rempli {len(filled)} paragraphe(s).")

    for para_id, filled_text in filled.items():
        try:
            idx = int(para_id)
        except ValueError:
            continue
        if idx >= len(paragraphs):
            continue

        para = paragraphs[idx]
        if para.text == filled_text:
            continue

        pairs = extract_replacement_pairs(para.text, filled_text)
        if pairs:
            for old_text, new_text in pairs:
                _replace_in_runs(para.runs, old_text, new_text)
            continue

        if para.runs:
            para.runs[0].text = filled_text
            for run in para.runs[1:]:
                run.text = ""

    doc.save(str(dst))


# ===========================================================================
# PIPELINE 3 : PDF SCANNÉ
# ===========================================================================

def _build_docx_from_paragraphs(paragraphs: list[dict[str, Any]], dst: Path) -> None:
    if not DOCX_AVAILABLE:
        sys.exit("Installer python-docx : pip install python-docx")

    from docx import Document as DocxDocument
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Pt, RGBColor

    doc = DocxDocument()
    for section in doc.sections:
        section.top_margin    = Pt(56)
        section.bottom_margin = Pt(56)
        section.left_margin   = Pt(70)
        section.right_margin  = Pt(70)

    style_map = {
        "title":    {"size": 12,  "bold": True,  "align": WD_ALIGN_PARAGRAPH.CENTER,  "sb": 4, "sa": 4},
        "heading":  {"size": 10,  "bold": True,  "align": WD_ALIGN_PARAGRAPH.LEFT,    "sb": 6, "sa": 2},
        "subhead":  {"size": 10,  "bold": True,  "align": WD_ALIGN_PARAGRAPH.LEFT,    "sb": 4, "sa": 1},
        "body":     {"size": 10,  "bold": False, "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "sb": 1, "sa": 1},
        "bullet":   {"size": 10,  "bold": False, "align": WD_ALIGN_PARAGRAPH.LEFT,    "sb": 1, "sa": 1},
        "footnote": {"size": 8,   "bold": False, "align": WD_ALIGN_PARAGRAPH.LEFT,    "sb": 1, "sa": 1},
    }

    for item in paragraphs:
        kind  = item.get("type", "body")
        text  = _pdf_safe_text(item["text"])
        st    = style_map.get(kind, style_map["body"])
        lines = text.split("\\n") if "\\n" in text else text.split("\n")
        lines = [ln for ln in lines if ln.strip()]

        for line_text in lines:
            para = doc.add_paragraph()
            para.alignment = st["align"]
            para.paragraph_format.space_before      = Pt(st["sb"])
            para.paragraph_format.space_after       = Pt(st["sa"])
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            if kind == "bullet":
                para.paragraph_format.left_indent = Pt(18)
            run = para.add_run(line_text.strip())
            run.bold           = st["bold"]
            run.font.name      = "Times New Roman"
            run.font.size      = Pt(st["size"])
            run.font.color.rgb = RGBColor(0, 0, 0)

    doc.save(str(dst))


def _find_system_fonts() -> tuple[str | None, str | None]:
    pairs = [
        (
            "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf",
        ),
        (
            "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf",
        ),
        (
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        ),
    ]
    for regular, bold in pairs:
        if Path(regular).exists() and Path(bold).exists():
            return regular, bold
    return None, None


_SYSTEM_FONT_PATH, _SYSTEM_FONT_BOLD_PATH = _find_system_fonts()


def _build_pdf_from_paragraphs(paragraphs: list[dict[str, Any]], dst: Path) -> None:
    x0     = PAGE_MARGIN_X_PT + 14
    x1     = PAGE_WIDTH_PT - PAGE_MARGIN_X_PT - 14
    text_w = x1 - x0

    style_map = {
        "title":    {"size": 12.0, "bold": True,  "align": fitz.TEXT_ALIGN_CENTER,  "space": 5},
        "heading":  {"size": 10.0, "bold": True,  "align": fitz.TEXT_ALIGN_LEFT,    "space": 4},
        "subhead":  {"size": 10.0, "bold": True,  "align": fitz.TEXT_ALIGN_LEFT,    "space": 3},
        "body":     {"size": 10.0, "bold": False, "align": fitz.TEXT_ALIGN_JUSTIFY, "space": 2},
        "bullet":   {"size": 10.0, "bold": False, "align": fitz.TEXT_ALIGN_LEFT,    "space": 2},
        "footnote": {"size": 8.0,  "bold": False, "align": fitz.TEXT_ALIGN_LEFT,    "space": 1},
    }

    doc  = fitz.open()
    page = doc.new_page(width=PAGE_WIDTH_PT, height=PAGE_HEIGHT_PT)
    y    = float(PAGE_MARGIN_TOP_PT)

    font_reg  = "helv"
    font_bold = "helv"
    if _SYSTEM_FONT_PATH and _SYSTEM_FONT_BOLD_PATH:
        try:
            page.insert_font(fontname="DocReg",  fontfile=_SYSTEM_FONT_PATH)
            page.insert_font(fontname="DocBold", fontfile=_SYSTEM_FONT_BOLD_PATH)
            font_reg  = "DocReg"
            font_bold = "DocBold"
        except Exception:
            pass

    def new_page() -> Any:
        nonlocal page, y
        page = doc.new_page(width=PAGE_WIDTH_PT, height=PAGE_HEIGHT_PT)
        y    = float(PAGE_MARGIN_TOP_PT)
        if font_reg != "helv":
            try:
                page.insert_font(fontname="DocReg",  fontfile=_SYSTEM_FONT_PATH)
                page.insert_font(fontname="DocBold", fontfile=_SYSTEM_FONT_BOLD_PATH)
            except Exception:
                pass

    for item in paragraphs:
        kind  = item.get("type", "body")
        text  = _pdf_safe_text(item["text"])
        st    = style_map.get(kind, style_map["body"])
        fname = font_bold if st["bold"] else font_reg

        bx0 = x0 + (18 if kind == "bullet" else 0)
        bw  = text_w - (18 if kind == "bullet" else 0)

        line_h         = st["size"] * 1.3
        chars_per_line = max(1, int(bw / (st["size"] * 0.55)))
        est_lines      = max(1, -(-len(text) // chars_per_line))
        block_h        = est_lines * line_h + st["space"]

        if y + block_h > PAGE_HEIGHT_PT - PAGE_MARGIN_BOT_PT and y > PAGE_MARGIN_TOP_PT + 10:
            new_page()

        rect = fitz.Rect(bx0, y, bx0 + bw, y + block_h + 16)
        res  = page.insert_textbox(
            rect, text,
            fontname=fname, fontsize=st["size"],
            color=(0, 0, 0), align=st["align"],
        )

        actual_h = block_h if res >= 0 else block_h * 2
        y += actual_h + st["space"]

        if y > PAGE_HEIGHT_PT - PAGE_MARGIN_BOT_PT:
            new_page()

    doc.save(str(dst), garbage=4, deflate=True)
    doc.close()


def convert_docx_to_pdf(docx_path: Path, pdf_path: Path) -> bool:
    office_bin = shutil.which("soffice") or shutil.which("libreoffice")
    if office_bin is None:
        return False
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        result = subprocess.run(
            [office_bin, "--headless", "--convert-to", "pdf",
             "--outdir", str(pdf_path.parent), str(docx_path)],
            check=False, capture_output=True, text=True,
        )
    except OSError:
        return False
    converted_path = pdf_path.parent / f"{docx_path.stem}.pdf"
    if result.returncode != 0 or not converted_path.exists():
        return False
    if converted_path != pdf_path:
        converted_path.replace(pdf_path)
    return True


def process_scanned_pdf(
    src: Path,
    dst: Path,
    api_key: str,
    doc_type: str = "unknown",
    pages: list[int] | None = None,
    company_info: dict[str, str] | None = None,
) -> None:
    try:
        from pdf2image import convert_from_path
    except ImportError:
        sys.exit("Installer les dépendances : pip install pdf2image Pillow")

    if pages is not None:
        first  = pages[0] + 1
        last   = pages[-1] + 1
        print(f"  Conversion des pages {first} à {last} en images ({SCAN_DPI} dpi)...")
        images = convert_from_path(str(src), dpi=SCAN_DPI, first_page=first, last_page=last)
    else:
        print(f"  Conversion de toutes les pages en images ({SCAN_DPI} dpi)...")
        images = convert_from_path(str(src), dpi=SCAN_DPI)

    print(f"  {len(images)} page(s) — envoi à Pixtral...")
    paragraphs = call_pixtral_vision(images, api_key, doc_type=doc_type, company_info=company_info)

    if not paragraphs:
        sys.exit("Pixtral a retourné un document vide.")

    print(f"  Pixtral a extrait {len(paragraphs)} paragraphe(s).")

    dst_docx = dst.with_suffix(".docx")
    _build_docx_from_paragraphs(paragraphs, dst_docx)

    if dst.suffix.lower() == ".pdf":
        if convert_docx_to_pdf(dst_docx, dst):
            print(f"  Conversion DOCX->PDF -> {dst}")
        else:
            print(f"  Reconstruction PDF  -> {dst} (fallback)")
            _build_pdf_from_paragraphs(paragraphs, dst)
