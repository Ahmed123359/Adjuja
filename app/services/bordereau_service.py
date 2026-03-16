import base64
import io
import json
from typing import Optional

import fitz  # pymupdf
from openai import AsyncOpenAI

from app.config.settings import get_settings

KEYWORDS = [
    "bordereau",
    "prix unitaire",
    "hors tva",
    "désignation",
    "désignations",
    "montant total",
    "unité de compte",
    "prix-detail",
    "détail estimatif",
    "forfait",
    "prix total",
]

SYSTEM_PROMPT = (
    "Tu es un assistant spécialisé dans l'analyse de documents d'appels d'offres marocains. "
    "Tu extrais avec précision les tableaux de bordereaux de prix. "
    "Réponds uniquement en JSON valide, sans markdown."
)

USER_PROMPT = """Cette page contient un tableau de bordereau de prix issu d'un appel d'offres.

Extrait le tableau tel qu'il apparaît dans le document, en t'adaptant aux colonnes réellement présentes.

Retourne un JSON avec cette structure :
{
  "titre": "titre exact du tableau tel qu'il apparaît dans le document",
  "colonnes": ["liste", "des", "colonnes", "telles", "quelles"],
  "lignes": [
    {"valeurs": ["valeur col1", "valeur col2", "..."]}
  ],
  "totaux": [
    {"label": "Montant Total Hors TVA", "valeur": null},
    {"label": "Taux de la TVA (20%)", "valeur": "20%"},
    {"label": "Montant Total TTC", "valeur": null}
  ]
}

Règles :
- Respecte l'ordre et le nombre exact de colonnes du document
- Les cellules vides (à remplir par le soumissionnaire) ont la valeur null
- Les cellules renseignées gardent leur valeur exacte
- Ne déduis rien, ne complète rien — recopie uniquement ce qui est visible"""


async def detect_and_extract_bordereau(pdf_bytes: bytes) -> Optional[dict]:
    """
    Étape 1 (gratuit) : détecte la page avec le plus de mots-clés de bordereau.
    Étape 2 (GPT-4o)  : extrait le tableau de cette page en JSON.
    Retourne None si aucune page candidate (score < 2).
    """
    settings = get_settings()

    # ── Étape 1 : détection par mots-clés ─────────────────────
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    best_index = None
    best_score = 0

    for i, page in enumerate(doc):
        text = page.get_text().lower()
        score = sum(1 for kw in KEYWORDS if kw in text)
        if score > best_score:
            best_score = score
            best_index = i

    if best_score < 2 or best_index is None:
        doc.close()
        return None

    # ── Étape 2 : rendu en image haute résolution ──────────────
    page = doc[best_index]
    mat = fitz.Matrix(2.0, 2.0)  # zoom x2 = 144 DPI
    pix = page.get_pixmap(matrix=mat)
    img_bytes = pix.tobytes("png")
    img_b64 = base64.b64encode(img_bytes).decode()
    doc.close()

    # ── Étape 3 : extraction GPT-4o vision ─────────────────────
    client = AsyncOpenAI(api_key=settings.openai_api_key)
    response = await client.chat.completions.create(
        model="gpt-4o",
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": USER_PROMPT},
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:image/png;base64,{img_b64}",
                            "detail": "high",
                        },
                    },
                ],
            },
        ],
        max_tokens=2000,
        temperature=0,
    )

    raw = response.choices[0].message.content.strip()
    if raw.startswith("```"):
        raw = raw.split("\n", 1)[1].rsplit("```", 1)[0]

    result = json.loads(raw)
    result["_page"] = best_index + 1
    return result


def bordereau_to_excel(bordereau: dict) -> bytes:
    """Convertit un dict bordereau (extrait par GPT-4o) en fichier Excel (.xlsx)."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bordereau de prix"

    style_titre = Font(bold=True, size=13)
    style_header = Font(bold=True, color="FFFFFF")
    fill_header = PatternFill("solid", fgColor="2E4057")
    fill_vide = PatternFill("solid", fgColor="FFF176")   # jaune = à remplir
    fill_total = PatternFill("solid", fgColor="E8F5E9")  # vert = totaux
    border_thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    row = 1
    ws.cell(row=row, column=1, value=bordereau.get("titre", "Bordereau de prix")).font = style_titre
    row += 2

    colonnes = bordereau.get("colonnes", [])
    for col_idx, col_name in enumerate(colonnes, start=1):
        cell = ws.cell(row=row, column=col_idx, value=col_name)
        cell.font      = style_header
        cell.fill      = fill_header
        cell.alignment = center
        cell.border    = border_thin
    row += 1

    for ligne in bordereau.get("lignes", []):
        for col_idx, valeur in enumerate(ligne.get("valeurs", []), start=1):
            cell = ws.cell(row=row, column=col_idx, value=valeur)
            cell.border    = border_thin
            cell.alignment = center if col_idx != 2 else left
            if valeur is None:
                cell.fill = fill_vide
        row += 1

    totaux = bordereau.get("totaux", [])
    if totaux:
        row += 1
        nb_cols = len(colonnes)
        for total in totaux:
            lc = ws.cell(row=row, column=max(1, nb_cols - 1), value=total.get("label"))
            lc.font = Font(bold=True)
            lc.fill = fill_total
            lc.border = border_thin
            lc.alignment = Alignment(horizontal="right")

            vc = ws.cell(row=row, column=max(1, nb_cols), value=total.get("valeur"))
            vc.fill   = fill_vide if total.get("valeur") is None else fill_total
            vc.border = border_thin
            vc.alignment = center
            row += 1

    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
