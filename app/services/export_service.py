"""
Export service  génère un vrai fichier .docx à partir d'un GenerationResult.

Convertit le Markdown des sections en styles Word natifs (titres, listes,
gras, italique) pour un rendu correct sur desktop et mobile.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor, Inches

from app.models.generation import GenerationResult


# ── Palette ────────────────────────────────────────────────────────────────
_BLUE  = RGBColor(0x1B, 0x3F, 0x6B)   # Navy bleu titres
_TEAL  = RGBColor(0x17, 0xA5, 0x89)   # Vert-bleu accents
_GREY  = RGBColor(0x6B, 0x72, 0x80)   # Gris métadonnées
_BLACK = RGBColor(0x1A, 0x1A, 0x2E)   # Corps de texte


# ── Helpers XML ────────────────────────────────────────────────────────────

def _set_cell_bg(cell, hex_color: str) -> None:
    """Colorise le fond d'une cellule de tableau."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _add_horizontal_rule(doc: Document, color: str = "D5E8F5") -> None:
    """Ligne de séparation horizontale."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


# ── Markdown inline parser ─────────────────────────────────────────────────

def _add_inline(run_container, text: str) -> None:
    """
    Découpe `text` selon les marqueurs inline Markdown et ajoute les runs
    correspondants avec les styles bold/italic.
    """
    # Pattern : **gras**, *italique*, `code`
    pattern = re.compile(r'\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`')
    pos = 0
    for m in pattern.finditer(text):
        # Texte avant le marqueur
        if m.start() > pos:
            run = run_container.add_run(text[pos:m.start()])
            run.font.color.rgb = _BLACK
        if m.group(1):   # **gras**
            run = run_container.add_run(m.group(1))
            run.bold = True
            run.font.color.rgb = _BLACK
        elif m.group(2): # *italique*
            run = run_container.add_run(m.group(2))
            run.italic = True
            run.font.color.rgb = _BLACK
        elif m.group(3): # `code`
            run = run_container.add_run(m.group(3))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xE8, 0x3E, 0x3E)
        pos = m.end()
    # Reste du texte
    if pos < len(text):
        run = run_container.add_run(text[pos:])
        run.font.color.rgb = _BLACK


# ── Section Markdown → paragraphes Word ────────────────────────────────────

def _render_section_content(doc: Document, markdown: str) -> None:
    """
    Parse le contenu Markdown d'une section et l'ajoute au document Word
    avec les styles appropriés.
    """
    lines = markdown.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Titre H2
        if line.startswith("## "):
            p = doc.add_paragraph()
            p.style = "Heading 2"
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after  = Pt(4)
            run = p.add_run(line[3:].strip())
            run.font.color.rgb = _BLUE
            run.font.size = Pt(11)
            run.bold = True

        # Titre H3
        elif line.startswith("### "):
            p = doc.add_paragraph()
            p.style = "Heading 3"
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after  = Pt(3)
            run = p.add_run(line[4:].strip())
            run.font.color.rgb = _TEAL
            run.font.size = Pt(10.5)
            run.bold = True

        # Titre H4
        elif line.startswith("#### "):
            p = doc.add_paragraph()
            run = p.add_run(line[5:].strip())
            run.bold = True
            run.italic = True
            run.font.color.rgb = RGBColor(0x24, 0x71, 0xA3)
            run.font.size = Pt(10)
            p.paragraph_format.space_before = Pt(6)

        # Liste à puces (- ou *)
        elif re.match(r'^[-*•]\s+', line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent    = Cm(0.8)
            p.paragraph_format.space_after    = Pt(2)
            p.paragraph_format.space_before   = Pt(1)
            text = re.sub(r'^[-*•]\s+', '', line)
            _add_inline(p, text)

        # Liste numérotée
        elif re.match(r'^\d+\.\s+', line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent    = Cm(0.8)
            p.paragraph_format.space_after    = Pt(2)
            text = re.sub(r'^\d+\.\s+', '', line)
            _add_inline(p, text)

        # Ligne vide → petit espace
        elif line.strip() == "":
            if i + 1 < len(lines) and lines[i + 1].strip() != "":
                doc.add_paragraph().paragraph_format.space_after = Pt(2)

        # Séparateur ---
        elif re.match(r'^---+$', line.strip()):
            _add_horizontal_rule(doc)

        # Paragraphe normal
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after  = Pt(5)
            p.paragraph_format.alignment    = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_inline(p, line)
            # Appliquer la police sur tous les runs qui n'en ont pas
            for run in p.runs:
                if not run.font.size:
                    run.font.size = Pt(10.5)

        i += 1


# ── Builder principal ──────────────────────────────────────────────────────

def build_docx(
    result: GenerationResult,
    company_nom: str = "",
    ao_text: str = "",
) -> bytes:
    """
    Génère un fichier .docx à partir d'un GenerationResult.
    Retourne les bytes du fichier prêt à être streamé.
    """
    doc = Document()

    # ── Marges de page ──────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # ── Styles de base ──────────────────────────────────────────────────
    normal = doc.styles["Normal"]
    normal.font.name  = "Calibri"
    normal.font.size  = Pt(10.5)
    normal.font.color.rgb = _BLACK

    # ── PAGE DE COUVERTURE ───────────────────────────────────────────────
    now     = datetime.now().strftime("%d/%m/%Y")
    ao_line = next((l.strip() for l in ao_text.split("\n") if len(l.strip()) > 10), "Appel d'offres")
    ao_line = ao_line[:130] + ("…" if len(ao_line) > 130 else "")

    # Titre principal
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(60)
    title_p.paragraph_format.space_after  = Pt(12)
    run = title_p.add_run("Réponse à l'Appel d'Offres")
    run.font.size  = Pt(24)
    run.font.bold  = True
    run.font.color.rgb = _BLUE

    # Ligne décorative
    _add_horizontal_rule(doc, "D5E8F5")

    # Intitulé de l'AO
    ao_p = doc.add_paragraph()
    ao_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ao_p.paragraph_format.space_before = Pt(14)
    ao_p.paragraph_format.space_after  = Pt(14)
    run = ao_p.add_run(ao_line)
    run.font.size  = Pt(13)
    run.font.bold  = True
    run.font.color.rgb = _BLACK

    _add_horizontal_rule(doc, "D5E8F5")

    # Infos entreprise + date
    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_p.paragraph_format.space_before = Pt(20)
    run_nom = info_p.add_run(company_nom or "Répondant")
    run_nom.font.size  = Pt(11)
    run_nom.font.bold  = True
    run_nom.font.color.rgb = _BLUE
    info_p.add_run(f"\n{now}").font.color.rgb = _GREY

    # Saut de page
    doc.add_page_break()

    # ── SOMMAIRE ────────────────────────────────────────────────────────
    toc_title = doc.add_paragraph()
    toc_title.paragraph_format.space_after = Pt(10)
    run = toc_title.add_run("Sommaire")
    run.font.size  = Pt(16)
    run.font.bold  = True
    run.font.color.rgb = _BLUE

    _add_horizontal_rule(doc)

    for s in result.sections:
        toc_p = doc.add_paragraph()
        toc_p.paragraph_format.space_after = Pt(3)
        toc_p.paragraph_format.left_indent = Cm(0.3)
        run_num = toc_p.add_run(f"{s.ordre + 1}.  ")
        run_num.font.bold  = True
        run_num.font.color.rgb = _TEAL
        run_title = toc_p.add_run(s.titre)
        run_title.font.color.rgb = _BLACK
        run_title.font.size = Pt(10.5)

    doc.add_page_break()

    # ── SECTIONS ────────────────────────────────────────────────────────
    for s in result.sections:
        # En-tête de section
        sec_title = doc.add_paragraph()
        sec_title.paragraph_format.space_before = Pt(4)
        sec_title.paragraph_format.space_after  = Pt(8)
        run = sec_title.add_run(f"{s.ordre + 1}.  {s.titre}")
        run.font.size  = Pt(14)
        run.font.bold  = True
        run.font.color.rgb = _BLUE

        _add_horizontal_rule(doc)

        # Contenu Markdown → Word
        _render_section_content(doc, s.contenu)

        doc.add_page_break()

    # ── Sérialisation ────────────────────────────────────────────────────
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
