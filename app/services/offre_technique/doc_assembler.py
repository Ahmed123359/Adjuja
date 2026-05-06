# -- coding: utf-8 --
import io
import json
import logging
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.models.offre_technique import CPSContext

logger = logging.getLogger(__name__)

# ── Palette par défaut ───────────────────────────────────────────────────────
_H1_COLOR     = "1F3864"
_H2_COLOR     = "2C3E50"
_BODY_COLOR   = "000000"
_TABLE_HEADER = "1F3864"
_FONT         = "Times New Roman"

SECTION_TITLES = {
    "methodologie": "1. Méthodologie et Organisation",
    "moyens":       "2. Moyens Humains et Matériels",
    "planning":     "3. Planning d'Exécution",
    "rse":          "4. Démarche RSE et Développement Durable",
    "references":   "5. Références Similaires et Fiches Techniques",
}

_BOLD_RE   = re.compile(r"\*\*(.+?)\*\*")
_ITALIC_RE = re.compile(r"\*([^*\n]+?)\*")

# Lines the AI sometimes emits that have no place in a formal document
_AI_NOISE_RE = re.compile(
    r"^("
    r"voici\s|cette section |ce planning |en résumé[,\s]|"
    r"notre (approche|méthodologie|engagement)|"
    r"note\s*:|nb\s*:|ps\s*:"
    r")",
    re.IGNORECASE,
)
# Parenthetical meta-notes like *(Sélection de 3 missions)* or *(Alignées sur les critères...)*
_META_NOTE_RE = re.compile(r"^\*?\([^)]{3,}\)\*?\.?$")


# ── XML helpers ──────────────────────────────────────────────────────────────

def _hex_to_rgb(h: str) -> tuple[int, int, int]:
    h = h.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _force_font(run, size_pt: float) -> None:
    """Force Times New Roman by writing XML directly (bypasses theme fonts)."""
    rpr = run._r.get_or_add_rPr()
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"),    _FONT)
    rfonts.set(qn("w:hAnsi"),    _FONT)
    rfonts.set(qn("w:cs"),       _FONT)
    rfonts.set(qn("w:eastAsia"), _FONT)
    rpr.insert(0, rfonts)
    run.font.size = Pt(size_pt)


def _set_doc_defaults(doc: Document) -> None:
    """Apply Times New Roman at the document default level."""
    styles_el = doc.styles.element
    doc_defaults = styles_el.find(qn("w:docDefaults"))
    if doc_defaults is None:
        return
    rpr_default = doc_defaults.find(f".//{qn('w:rPrDefault')}/{qn('w:rPr')}")
    if rpr_default is None:
        return
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"),    _FONT)
    rfonts.set(qn("w:hAnsi"),    _FONT)
    rfonts.set(qn("w:cs"),       _FONT)
    rfonts.set(qn("w:eastAsia"), _FONT)
    rpr_default.insert(0, rfonts)


def _para_format(para, space_after: int = 6, line_spacing: float = 1.15) -> None:
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = para.paragraph_format
    fmt.space_after  = Pt(space_after)
    fmt.line_spacing = Pt(11 * line_spacing)


def _strip_md(text: str) -> str:
    return _BOLD_RE.sub(r"\1", text).replace("*", "").strip()


def _add_rich_paragraph(doc: Document, text: str, style: str | None = None,
                         size: float = 11, color_hex: str = _BODY_COLOR) -> None:
    """Paragraph with inline **bold** and proper formatting."""
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    _para_format(p)
    text = re.sub(r"^[✅✔☑►•·]\s*", "", text.strip())
    text = _ITALIC_RE.sub(r"**\1**", text)
    parts = _BOLD_RE.split(text)
    r, g, b = _hex_to_rgb(color_hex)
    for i, part in enumerate(parts):
        if not part:
            continue
        run = p.add_run(part.replace("*", ""))   # strip any remaining lone *
        run.bold = bool(i % 2)
        run.font.color.rgb = RGBColor(r, g, b)
        _force_font(run, size)


def _add_heading(doc: Document, text: str, level: int, color_hex: str) -> None:
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    size = 13 if level == 1 else 12 if level == 2 else 11
    r, g, b = _hex_to_rgb(color_hex)
    for run in p.runs:
        run.font.color.rgb = RGBColor(r, g, b)
        _force_font(run, size)


def _add_markdown_table(doc: Document, lines: list[str], header_color: str) -> None:
    rows = []
    for line in lines:
        stripped = line.strip().strip("|")
        if re.match(r"^[\s\-|:]+$", stripped):
            continue
        cells = [c.strip() for c in stripped.split("|")]
        if any(cells):
            rows.append(cells)
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    r, g, b = _hex_to_rgb(header_color)
    for i, row in enumerate(rows):
        for j in range(cols):
            cell_text = row[j] if j < len(row) else ""
            cell = table.rows[i].cells[j]
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run  = para.add_run(_strip_md(cell_text))
            _force_font(run, 10)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                tc_pr = cell._tc.get_or_add_tcPr()
                shd   = OxmlElement("w:shd")
                shd.set(qn("w:val"),   "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"),  header_color)
                tc_pr.append(shd)
    doc.add_paragraph()


# ── Gantt via matplotlib ─────────────────────────────────────────────────────

def _build_gantt(phases: list[dict], brand_color: str) -> bytes | None:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        import matplotlib.patches as mpatches
        import numpy as np

        n      = len(phases)
        fig, ax = plt.subplots(figsize=(12, max(3, n * 0.6 + 1)))
        ax.set_facecolor("#FAFAFA")
        fig.patch.set_facecolor("white")

        main_color = f"#{brand_color.lstrip('#')}" if brand_color else "#1F3864"
        alt_color  = "#2C3E50"

        for i, phase in enumerate(phases):
            color  = main_color if i % 2 == 0 else alt_color
            label  = phase.get("nom", f"Phase {i+1}")
            duree  = phase.get("duree", "")
            ax.barh(i, 1, left=i, color=color, edgecolor="white", height=0.6)
            ax.text(i + 0.5, i, f"{label}\n{duree}", va="center", ha="center",
                    fontsize=8, color="white", fontweight="bold",
                    fontfamily="DejaVu Sans")

        ax.set_yticks(range(n))
        ax.set_yticklabels([p.get("nom", "") for p in phases], fontsize=8)
        ax.set_xticks([])
        ax.invert_yaxis()
        ax.set_title("Planning d'exécution", fontsize=10, fontweight="bold",
                      color="#1F3864", fontfamily="DejaVu Sans")
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        ax.spines["bottom"].set_visible(False)

        buf = io.BytesIO()
        plt.savefig(buf, format="png", dpi=150, bbox_inches="tight")
        plt.close(fig)
        buf.seek(0)
        return buf.read()
    except Exception as exc:
        logger.warning("Gantt generation failed: %s", exc)
        return None


# ── Planning section ─────────────────────────────────────────────────────────

_DEFAULT_PHASES = [
    {"nom": "Phase 1 : Préparation",       "duree": "2 semaines"},
    {"nom": "Phase 2 : Lancement",         "duree": "2 semaines"},
    {"nom": "Phase 3 : Exécution",         "duree": "8 semaines"},
    {"nom": "Phase 4 : Suivi & contrôle",  "duree": "4 semaines"},
    {"nom": "Phase 5 : Clôture",           "duree": "2 semaines"},
]


def _add_planning_table(doc: Document, planning_text: str,
                         header_color: str, brand_color: str | None) -> None:
    try:
        clean = re.sub(r"```(?:json)?\s*", "", planning_text, flags=re.IGNORECASE).strip()
        clean = re.sub(r"```\s*$", "", clean, flags=re.IGNORECASE).strip()
        m     = re.search(r"\{.*\}", clean, re.DOTALL)
        if not m:
            gantt_png = _build_gantt(_DEFAULT_PHASES, brand_color or "")
            if gantt_png:
                doc.add_picture(io.BytesIO(gantt_png), width=Inches(6))
                cap = doc.add_paragraph("Figure 1 : Planning d'exécution")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _render_section(doc, planning_text, header_color)
            return

        data = json.loads(m.group(0))
        if "phases" not in data:
            for v in data.values():
                if isinstance(v, dict) and "phases" in v:
                    data = v
                    break

        phases    = data.get("phases", []) or _DEFAULT_PHASES
        narrative = data.get("description", "")

        # Try Gantt chart first
        if phases:
            gantt_png = _build_gantt(phases, brand_color or "")
            if gantt_png:
                doc.add_picture(io.BytesIO(gantt_png), width=Inches(6))
                cap = doc.add_paragraph("Figure 1 : Planning d'exécution")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in cap.runs:
                    run.font.size = Pt(9)
                    run.font.italic = True
                    _force_font(run, 9)
                doc.add_paragraph()
            else:
                # Fallback: Word table
                table = doc.add_table(rows=1, cols=4)
                table.style = "Table Grid"
                hdr = table.rows[0].cells
                r, g, b = _hex_to_rgb(header_color)
                for i, h in enumerate(["Phase", "Durée", "Jalons", "Livrables"]):
                    para = hdr[i].paragraphs[0]
                    run  = para.add_run(h)
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    _force_font(run, 10)
                    tc_pr = hdr[i]._tc.get_or_add_tcPr()
                    shd   = OxmlElement("w:shd")
                    shd.set(qn("w:val"),   "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"),  header_color)
                    tc_pr.append(shd)
                for phase in phases:
                    row = table.add_row().cells
                    row[0].text = phase.get("nom", "")
                    row[1].text = phase.get("duree", "")
                    row[2].text = "\n".join(phase.get("jalons", []))
                    row[3].text = "\n".join(phase.get("livrables", []))
                doc.add_paragraph()

        if narrative:
            _add_rich_paragraph(doc, narrative)

    except (json.JSONDecodeError, KeyError):
        _render_section(doc, planning_text, header_color)


# ── Section renderer ─────────────────────────────────────────────────────────

def _render_section(doc: Document, text: str, h1_color: str) -> None:
    lines     = text.split("\n")
    in_code   = False
    table_buf: list[str] = []
    first_heading_skipped = False

    def flush_table() -> None:
        if table_buf:
            _add_markdown_table(doc, table_buf, h1_color)
            table_buf.clear()

    for line in lines:
        stripped = line.strip()

        if stripped.startswith("```"):
            in_code = not in_code
            flush_table()
            continue
        if in_code:
            continue
        if not stripped:
            flush_table()
            continue
        if stripped.startswith("|"):
            table_buf.append(stripped)
            continue
        else:
            flush_table()

        if re.match(r"^[-*_]{3,}$", stripped):
            continue

        # Drop AI meta-commentary and parenthetical notes
        if _AI_NOISE_RE.search(stripped) or _META_NOTE_RE.match(stripped):
            continue

        heading_m = re.match(r"^(#{1,4})\s+(.+)", stripped)
        if heading_m:
            if not first_heading_skipped:
                first_heading_skipped = True
                continue
            depth  = len(heading_m.group(1))
            title  = _strip_md(heading_m.group(2))
            color  = h1_color if depth <= 2 else _H2_COLOR
            level  = min(depth + 1, 4)
            _add_heading(doc, title, level, color)
            continue

        if re.match(r"^\*\*.+\*\*$", stripped) and len(stripped) < 80:
            _add_heading(doc, _strip_md(stripped), 3, _H2_COLOR)
            continue

        if re.match(r"^[-*✅✔►]\s+", stripped):
            content = re.sub(r"^[-*✅✔►]\s+", "", stripped)
            _add_rich_paragraph(doc, content, style="List Bullet")
            continue

        _add_rich_paragraph(doc, stripped)

    flush_table()


# ── Cover page ───────────────────────────────────────────────────────────────

def _add_cover(doc: Document, cps: CPSContext, company_name: str,
               logo_bytes: bytes | None, brand_color: str | None) -> None:
    h_color = brand_color or f"#{_H1_COLOR}"
    r, g, b = _hex_to_rgb(h_color.lstrip("#"))

    # Logo top-right
    if logo_bytes:
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run()
            run.add_picture(io.BytesIO(logo_bytes), height=Inches(0.7))
        except Exception as exc:
            logger.warning("Logo insertion failed: %s", exc)

    doc.add_paragraph()

    title = doc.add_heading("OFFRE TECHNIQUE", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(r, g, b)
        _force_font(run, 22)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(f"Appel d'offres : {cps.reference or 'N/A'}")
    run.font.color.rgb = RGBColor(r, g, b)
    _force_font(run, 14)

    if cps.acheteur:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Acheteur : {cps.acheteur}")
        _force_font(run, 12)

    if company_name:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Soumissionnaire : {company_name}")
        run.bold = True
        _force_font(run, 12)

    doc.add_page_break()


# ── Main build ───────────────────────────────────────────────────────────────

def build_docx(
    sections:     dict[str, str],
    cps:          CPSContext,
    company_info: dict,
    output_path:  Path,
    logo_bytes:   bytes | None = None,
    logo_filename: str | None  = None,
    brand_color:  str | None   = None,
) -> Path:
    doc          = Document()
    company_name = company_info.get("company_name", "")
    h1_color     = brand_color.lstrip("#") if brand_color else _H1_COLOR

    _set_doc_defaults(doc)
    _add_cover(doc, cps, company_name, logo_bytes, brand_color)

    for key, title in SECTION_TITLES.items():
        _add_heading(doc, title, level=1, color_hex=h1_color)
        text = sections.get(key, "")

        if key == "planning":
            _add_planning_table(doc, text, h1_color, brand_color)
        else:
            _render_section(doc, text, h1_color)

        doc.add_page_break()

    doc.save(str(output_path))
    return output_path


def build_pdf(docx_path: Path) -> Path | None:
    pdf_path = docx_path.with_suffix(".pdf")
    try:
        result = subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf",
             "--outdir", str(docx_path.parent), str(docx_path)],
            capture_output=True, timeout=60,
        )
        if result.returncode == 0 and pdf_path.exists():
            return pdf_path
        logger.warning("LibreOffice conversion failed: %s", result.stderr.decode())
        return None
    except (FileNotFoundError, subprocess.TimeoutExpired) as e:
        logger.info("LibreOffice non disponible : %s", e)
        return None
